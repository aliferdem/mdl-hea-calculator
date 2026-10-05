"""Composition conversions, set-equal, normalise, sweep and expression generation, TSV, spreadsheets."""

import math

import pytest

from hea_calc.core import composition as C
from hea_calc.core.engine import Engine
from hea_calc.core.expression import ExpressionError, VariableRange, generate, parse_expression
from hea_calc.core.generation import ElementRange, RangeError, generate_sweep
from hea_calc.core.results import ResultRow, format_number, table_tsv
from hea_calc.core.spreadsheet import SpreadsheetError, export_rows, read_compositions


# ---------------------------------------------------------------- formulas
@pytest.mark.parametrize("text,expected", [
    ("CoCrFeNi", {"Co": 1, "Cr": 1, "Fe": 1, "Ni": 1}),
    ("Al0.5CoCrFeNi", {"Al": 0.5, "Co": 1, "Cr": 1, "Fe": 1, "Ni": 1}),
    ("(FeCoNi)90Al10", {"Fe": 30, "Co": 30, "Ni": 30, "Al": 10}),
    ("Al20(TiCo)80", {"Al": 20, "Ti": 40, "Co": 40}),
    ("Fe2Fe", {"Fe": 3}),
    ("W27.3Nb22.7Mo25.6Ta24.4", {"W": 27.3, "Nb": 22.7, "Mo": 25.6, "Ta": 24.4}),
])
def test_parse_formula(text, expected):
    got = C.parse_formula(text)
    assert got.keys() == expected.keys()
    for k in expected:
        assert got[k] == pytest.approx(expected[k])


@pytest.mark.parametrize("bad", ["", "Xx2", "(FeCo", "FeCo)", "Fe-Co", "fe"])
def test_parse_formula_errors(bad):
    with pytest.raises(C.FormulaError):
        C.parse_formula(bad)


# ---------------------------------------------------------------- conversions
def test_representation_round_trips():
    at = {"Al": 10.0, "Co": 22.5, "Cr": 22.5, "Fe": 22.5, "Ni": 22.5}
    wt = C.convert(at, "at", "wt")
    assert sum(wt.values()) == pytest.approx(100)
    assert wt["Al"] < 10  # Al is light
    back = C.convert(wt, "wt", "at")
    for el in at:
        assert back[el] == pytest.approx(at[el], rel=1e-12)
    mass = C.from_fractions(C.to_fractions(at, "at"), "mass", total_mass=2.0)
    assert sum(mass.values()) == pytest.approx(2.0)
    back = C.convert(mass, "mass", "at")
    for el in at:
        assert back[el] == pytest.approx(at[el], rel=1e-12)
    ratio = C.convert(at, "at", "ratio")
    assert ratio["Co"] == pytest.approx(1.0) and ratio["Al"] == pytest.approx(10 / 22.5)


def test_all_representations_keep_entered_column():
    vals = {"Fe": 1.0, "Ni": 3.0}
    reps = C.all_representations(vals, "ratio")
    assert reps["ratio"] == vals
    assert reps["at"] == pytest.approx({"Fe": 25.0, "Ni": 75.0})
    mass_in = {"Fe": 0.5, "Ni": 1.5}
    reps = C.all_representations(mass_in, "mass")
    assert sum(reps["mass"].values()) == pytest.approx(2.0)
    assert sum(reps["wt"].values()) == pytest.approx(100.0)


def test_conversion_rejects_invalid():
    with pytest.raises(ValueError):
        C.to_fractions({"Fe": -1, "Ni": 2}, "at")
    with pytest.raises(ValueError):
        C.to_fractions({"Fe": 0, "Ni": 0}, "wt")


def test_set_equal_respects_representation():
    els = ["Fe", "Co", "Ni", "Cr"]
    assert C.set_equal(els, "ratio") == {el: 1.0 for el in els}
    assert C.set_equal(els, "at") == {el: 25.0 for el in els}
    assert C.set_equal(els, "wt") == {el: 25.0 for el in els}
    assert C.set_equal(els, "mass", total_mass=2.0) == {el: 0.5 for el in els}
    with pytest.raises(ValueError):
        C.set_equal(els, "element")


def test_normalise():
    out = C.normalise({"Fe": 30.0, "Ni": 30.0}, "at")
    assert out == {"Fe": 50.0, "Ni": 50.0}
    assert C.can_normalise("wt") and not C.can_normalise("ratio") and not C.can_normalise("mass")
    for bad in ({"Fe": 0.0, "Ni": 0.0}, {"Fe": None, "Ni": 1.0}, {"Fe": -1.0, "Ni": 2.0},
                {"Fe": float("nan"), "Ni": 1.0}):
        with pytest.raises(ValueError):
            C.normalise(bad, "at")
    with pytest.raises(ValueError):
        C.normalise({"Fe": 1.0}, "ratio")


def test_parse_number():
    assert C.parse_number("1,5") == 1.5
    assert C.parse_number(" ") is None and C.parse_number("abc") is None and C.parse_number("inf") is None


def test_alloy_name():
    assert C.alloy_name({"Al": 0.5, "Co": 1, "Cr": 1, "Fe": 1, "Ni": 1}) == "Al0.5CoCrFeNi"
    assert C.alloy_name({"Co": 25, "Cr": 25, "Fe": 25, "Ni": 25}) == "CoCrFeNi"
    assert C.alloy_name({"Fe": 70, "Cr": 20, "Ni": 10}) == "FeCr0.286Ni0.143"
    assert C.name_html("Al0.5CoCrFeNi") == "Al<sub>0.5</sub>CoCrFeNi"


# ---------------------------------------------------------------- sweep
def test_sweep_ternary_count_and_sum():
    rs = [ElementRange(el, 0, 100, 10) for el in ("Al", "Ti", "V")]
    comps = generate_sweep(rs)
    assert len(comps) == 66  # C(12, 2)
    assert all(sum(c.values()) == pytest.approx(100) for c in comps)


def test_sweep_fixed_and_decimal_steps():
    rs = [ElementRange("Fe", 20, 20, 0), ElementRange("Co", 0.5, 79.5, 0.5), ElementRange("Ni", 0.5, 79.5, 0.5)]
    comps = generate_sweep(rs)
    assert len(comps) == 159
    assert all(c["Fe"] == 20 for c in comps)
    assert all(abs(sum(c.values()) - 100) < 1e-9 for c in comps)


def test_sweep_min_elements_and_percent_names():
    rs = [ElementRange(el, 0, 100, 50) for el in ("Al", "Ti")]
    assert generate_sweep(rs) == [{"Al": 0.0, "Ti": 100.0}, {"Al": 50.0, "Ti": 50.0}, {"Al": 100.0, "Ti": 0.0}]
    assert generate_sweep(rs, min_elements=2) == [{"Al": 50.0, "Ti": 50.0}]
    assert C.percent_name({"Al": 10.0, "Ti": 30.0, "V": 60.0, "Cr": 0.0}) == "Al10Ti30V60"
    assert C.percent_name({"Fe": 33.335, "Ni": 66.665}) == "Fe33.34Ni66.67" or C.percent_name(
        {"Fe": 33.335, "Ni": 66.665}).startswith("Fe33.3")


def test_sweep_balance_grid_and_errors():
    comps = generate_sweep([ElementRange("Fe", 0, 50, 25), ElementRange("Ni", 50, 100, 30)])
    assert comps == [{"Fe": 50.0, "Ni": 50.0}]  # Ni grid is 50, 80: balances 100 and 75 are rejected
    with pytest.raises(RangeError):
        generate_sweep([ElementRange("Fe", 0, 100, 10)])
    with pytest.raises(RangeError):
        generate_sweep([ElementRange("Fe", 10, 0, 1), ElementRange("Ni", 0, 100, 1)])
    with pytest.raises(RangeError):
        generate_sweep([ElementRange("Fe", 0, 100, 0), ElementRange("Ni", 0, 100, 1)])
    with pytest.raises(RangeError):
        generate_sweep([ElementRange(e, 0, 100, 1) for e in ("Fe", "Co", "Ni", "Cr")], limit=1000)


# ---------------------------------------------------------------- expression
def test_expression_variables_and_evaluation():
    ex = parse_expression("(FeCoNi)1-xAlx")
    assert ex.variables == ["x"]
    amounts = ex.evaluate({"x": 0.1})
    assert amounts["Al"] == pytest.approx(0.1)
    assert amounts["Fe"] == pytest.approx(0.3)
    assert ex.name({"x": 0.1}) == "(FeCoNi)0.9Al0.1"


@pytest.mark.parametrize("text,vars_,env,check", [
    ("Al_x(CoCrFeNi)", ["x"], {"x": 0.5}, {"Al": 0.5, "Co": 0.25}),
    ("Fe{2-x}Co{x}Ni", ["x"], {"x": 0.5}, {"Fe": 1.5, "Co": 0.5, "Ni": 1}),
    ("CoCrFeNiTi{0.1*k}", ["k"], {"k": 3}, {"Ti": 0.3}),
    ("(FeCoNi)100-cAlc", ["c"], {"c": 10}, {"Al": 10, "Fe": 30}),
    ("Fe{a}Ni{b}", ["a", "b"], {"a": 1, "b": 2}, {"Fe": 1, "Ni": 2}),
    ("Al2xNi", ["x"], {"x": 0.25}, {"Al": 0.5}),
    ("Fe_(1-y)^1Ni_y", None, None, None),
])
def test_expression_forms(text, vars_, env, check):
    if vars_ is None:
        with pytest.raises(ExpressionError):
            parse_expression(text)
        return
    ex = parse_expression(text)
    assert ex.variables == vars_
    got = ex.evaluate(env)
    for el, v in check.items():
        assert got[el] == pytest.approx(v)


def test_expression_precedence_in_simple_chain():
    ex = parse_expression("Fe1-2xNix")
    assert ex.evaluate({"x": 0.1})["Fe"] == pytest.approx(0.8)


@pytest.mark.parametrize("bad", ["", "Xy2", "(FeCo", "Fe{1-x", "Fe$"])
def test_expression_errors(bad):
    with pytest.raises(ExpressionError):
        parse_expression(bad)


def test_expression_generation_skips_invalid_and_duplicates():
    ex = parse_expression("(FeCoNi)1-xAlx")
    batch = generate(ex, {"x": VariableRange("x", 0, 1.2, 0.1)})
    # x = 0 -> FeCoNi (valid), x = 1 -> pure Al (skipped), x > 1 -> negative (skipped)
    assert len(batch.compositions) == 10
    assert batch.skipped == 3
    assert batch.names[0] == "FeCoNi"
    ex2 = parse_expression("Fe{x}Ni{x}")
    b2 = generate(ex2, {"x": VariableRange("x", 1, 3, 1)})
    assert len(b2.compositions) == 1 and b2.duplicates == 2
    with pytest.raises(RangeError):
        generate(ex, {})


def test_expression_feeds_shared_engine():
    eng = Engine()
    ex = parse_expression("Al{x}CoCrFeNi")
    batch = generate(ex, {"x": VariableRange("x", 0.5, 0.5, 0)})
    direct = eng.calculate({"Al": 0.5, "Co": 1, "Cr": 1, "Fe": 1, "Ni": 1})
    via = eng.calculate(batch.compositions[0])
    assert via.descriptors["dhmix"] == pytest.approx(direct.descriptors["dhmix"])


# ---------------------------------------------------------------- tables / clipboard / spreadsheets
def _rows():
    eng = Engine()
    out = []
    for comp in ({"Co": 1, "Cr": 1, "Fe": 1, "Ni": 1}, {"Al": 1, "Co": 1, "Cr": 1, "Fe": 1, "Ni": 1}):
        out.append(ResultRow(C.alloy_name(comp), comp, eng.calculate(comp)))
    return out


def test_table_tsv_for_excel():
    rows = _rows()
    text = table_tsv(rows, ["name", "dhmix", "delta", "M1", "M6"])
    lines = text.split("\r\n")
    assert lines[0] == "Alloy\tΔHmix (kJ/mol)\tδ (%)\tM1\tM6"
    assert lines[1].startswith("CoCrFeNi\t-3.75\t")
    assert lines[2].split("\t")[1] == "-12.32"
    assert "−" not in text and lines[-1] == ""
    assert all(len(l.split("\t")) == 5 for l in lines[:-1])


def test_format_number():
    assert format_number(-3.75) == "−3.75"
    assert format_number(-0.001) == "0.00"
    assert format_number(math.inf) == "∞"
    assert format_number(None) == "N/A"


def test_export_import_round_trip(tmp_path):
    rows = _rows()
    xlsx = export_rows(rows, tmp_path / "out.xlsx")
    import openpyxl
    ws = openpyxl.load_workbook(xlsx).worksheets[0]
    header = [c.value for c in ws[1]]
    i = header.index("ΔHmix (kJ/mol)")
    assert ws.cell(row=2, column=i + 1).value == pytest.approx(rows[0].result.descriptors["dhmix"], rel=1e-15)
    csvp = export_rows(rows, tmp_path / "out.csv")
    assert csvp.read_text(encoding="utf-8-sig").splitlines()[0].startswith("Alloy,Source")


def test_import_formula_column(tmp_path):
    p = tmp_path / "in.csv"
    p.write_text("Formula,Note\nAl0.5CoCrFeNi,a\nXy2Fe,b\n,c\nFe,d\n(FeCoNi)90Al10,e\n", encoding="utf-8")
    res = read_compositions(p)
    assert [e[0] for e in res.entries] == ["Al0.5CoCrFeNi", "(FeCoNi)90Al10"]
    rows = {n for n, _ in res.errors}
    assert rows == {3, 4, 5}


def test_import_element_columns_with_errors(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Fe", "Co (at.%)", "Ni", "Qq"])
    ws.append(["A", 25, 25, 50, 1])
    ws.append(["B", 30, "abc", 70, None])
    ws.append(["C", 50, None, None, None])
    ws.append(["D", -5, 50, 55, None])
    ws.append(["E", "33,3", 33.3, 33.4, None])
    p = tmp_path / "in.xlsx"
    wb.save(p)
    res = read_compositions(p)
    assert res.layout == "element columns"
    names = [n for n, _ in res.entries]
    assert names == ["A", "E"]
    assert res.entries[1][1]["Fe"] == pytest.approx(33.3)
    errs = dict(res.errors)
    assert 1 in errs and "Qq" in errs[1]  # invalid element column reported
    assert "not a number" in errs[3] and "Fewer than two" in errs[4] and "non-negative" in errs[5]


def test_import_without_composition_columns(tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(SpreadsheetError):
        read_compositions(p)


# ---------------------------------------------------------------- composition space (sweep design)
from hea_calc.core.generation import SweepSpec, generate_space, summarize_space, systems  # noqa: E402


def test_space_ratio_product_matches_figma_example():
    # Figma 278:17008: Cr 0.5-1.5/0.1, Mn 0.5-1.5/0.1, Fe 0.8-1.2/0.1, Co 1.0, Ni 0.5-1.5/0.1 -> 6,655
    specs = [SweepSpec("Cr", 0.5, 1.5, 0.1), SweepSpec("Mn", 0.5, 1.5, 0.1), SweepSpec("Fe", 0.8, 1.2, 0.1),
             SweepSpec("Co", 1.0, 1.0, 0.1), SweepSpec("Ni", 0.5, 1.5, 0.1)]
    s = summarize_space(specs)
    assert (s.selected, s.required, s.optional, s.systems, s.compositions) == (5, 5, 0, 1, 6655)


def test_space_optional_elements_and_system_sizes():
    specs = [SweepSpec(e, 1, 1, 0.1) for e in ("Fe", "Co", "Ni")] + \
            [SweepSpec(e, 1, 1, 0.1, required=False) for e in ("Cr", "Mn", "Al", "Ti")]
    assert len(systems(specs)) == 16  # all subsets of 4 optional elements
    assert len(systems(specs, {4, 5, 6, 7})) == 15  # Figma: "4 optional · 15 combinations"
    comps = generate_space(specs, sizes={4})
    assert len(comps) == 4 and all(len(c) == 4 for c in comps)


def test_space_removes_proportional_duplicates_and_at_percent_mode():
    specs = [SweepSpec("Fe", 1, 2, 1), SweepSpec("Ni", 1, 2, 1)]
    comps = generate_space(specs)
    assert len(comps) == 3  # (1,1) == (2,2)
    at = [SweepSpec("Fe", 10, 90, 10), SweepSpec("Ni", 10, 90, 10)]
    comps = generate_space(at, unit="at")
    assert len(comps) == 9 and all(sum(c.values()) == 100 for c in comps)
    zero_min = [SweepSpec("Fe", 0, 1, 0.5), SweepSpec("Ni", 0, 1, 0.5)]
    assert all(min(c.values()) > 0 for c in generate_space(zero_min))  # required -> present


def test_expression_blocks():
    from hea_calc.core.expression import block_variables, generate_blocks, parse_blocks, split_blocks
    text = "(FeCoNi)1-xAlx\n\nCoCrFeNi\nTi{y}\n\n\n(FeCoNi)1-xAlx"
    assert split_blocks(text) == ["(FeCoNi)1-xAlx", "CoCrFeNiTi{y}", "(FeCoNi)1-xAlx"]
    exprs = parse_blocks(text)
    assert block_variables(exprs) == ["x", "y"]
    b = generate_blocks(exprs, {"x": VariableRange("x", 0.1, 0.2, 0.1), "y": VariableRange("y", 0.5, 0.5, 0)})
    assert b.names == ["(FeCoNi)0.9Al0.1", "(FeCoNi)0.8Al0.2", "CoCrFeNiTi0.5"]
    assert b.duplicates == 2  # third block repeats the first
    with pytest.raises(ExpressionError, match="Block 2"):
        parse_blocks("FeCo\n\nXq2")
