"""SQLite result repository: identity, full-dataset sort/filter, N/A semantics, deletion, export,
single-writer behaviour, failures, provenance for future plots, lifecycle - and the table/Details
interactions built on it."""

import csv
import gc
import os
import sqlite3
import sys
import tempfile
import threading
import time
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QEvent, QObject, QPointF, QSettings, Qt  # noqa: E402
from PySide6.QtGui import QMouseEvent  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from hea_calc.core import repository as R  # noqa: E402
from hea_calc.core.chunks import pack_results, with_metadata  # noqa: E402
from hea_calc.core.engine import NA  # noqa: E402
from hea_calc.core.expression import VariableRange, generate_blocks, parse_blocks  # noqa: E402
from hea_calc.core.generation import SweepSpec  # noqa: E402
from hea_calc.core.jobs import CalculationJob, CalculationService, Delivery, JobInput  # noqa: E402
from hea_calc.core.recipes import ExpressionRecipe, RecalcRecipe, SweepRecipe  # noqa: E402
from hea_calc.core.repository import (ResultRepository, ResultView, cleanup_stale_sessions,  # noqa: E402
                                      filter_sql, order_sql)
from hea_calc.core.results import ModelFilter, RangeFilter, ResultRow  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
from synthetic import fill, synthetic_chunk  # noqa: E402


# =========================================================================== fixtures
@pytest.fixture()
def repo(engine, tmp_path):
    r = ResultRepository(engine, root=tmp_path / "sessions")
    yield r
    r.close()


def delivery(repo, run, svc=None):
    """Chunks of a job go to run ``run`` of the repository's writer process (as in the app)."""
    w = repo.writer
    if svc is not None and svc.channel is None and not svc.warm:
        svc.channel = w.inbox
    return Delivery(run, lambda i, ch: w.put_chunk(run, i, ch), lambda: w.stored_chunks(run),
                    lambda: bool(w.error), w.inbox)


def mixed_rows(engine, n=60):
    """Real results with ties, N/A values (an error row) and a spread of labels."""
    comps = [{"Fe": 1 + i % 7, "Co": 1 + i % 2, "Ni": 1 + i % 3, "Cr": 0.5 + i % 5, "Al": (i % 4) * 0.3}
             for i in range(n)]
    comps[7] = {"Fe": 1.0}                                        # engine error: every value N/A
    comps[13] = {"Fe": 1.0, "Xe": 1.0} if False else comps[13]
    return [ResultRow(f"R{i:03d}", c, engine.calculate(c), "sweep", {}) for i, c in enumerate(comps)]


def store_rows(repo, rows, mode="sweep"):
    run = repo.create_run(mode, settings={}, elements=["Fe", "Co", "Ni", "Cr", "Al"])
    rng = repo.append_results(run, rows)
    repo.finalize_run(run, "complete")
    return run, rng


def python_sort(rows, key, desc):
    """The in-memory semantics the SQL order must reproduce (former ResultStore.sort_ids)."""
    def val(r):
        return r.result.value(key) if r.result.ok else None
    present = [r for r in rows if val(r) is not None]
    missing = [r for r in rows if val(r) is None]
    present.sort(key=val, reverse=desc)
    return present + missing


# =========================================================================== identity / paging
def test_result_ids_are_stable_across_pages_sorts_and_filters(repo, engine):
    view = ResultView(repo)
    run = fill(repo, 250, engine)
    view.refresh()
    seen = {}
    for p in range(3):
        for rec in view.page(p * 100, 100):
            assert rec.id not in seen
            seen[rec.id] = rec.name
    assert len(seen) == 250 and list(seen) == sorted(seen), "natural order is id order"
    view.request(sort=("dhmix", True))
    for rec in view.page(100, 100):
        assert repo.get_result(rec.id).name == seen[rec.id] == rec.name
    view.request(filters=[RangeFilter("dhmix", None, -2.0)])
    for rec in view.page(0, 100):
        assert seen[rec.id] == rec.name
    assert repo.get_run(run).calculated == 250


# =========================================================================== full-dataset sort / filter
@pytest.mark.parametrize("key", ["dhmix", "omega", "name", "composition", "M1", "M6", "cstr"])
@pytest.mark.parametrize("desc", [False, True])
def test_sort_orders_the_complete_dataset_with_na_last(repo, engine, key, desc):
    rows = mixed_rows(engine)
    view = ResultView(repo)
    store_rows(repo, rows)
    view.refresh()
    view.request(sort=(key, desc))
    got = [r.name for r in view.iterate()]
    if key in ("name", "composition"):
        val = (lambda r: r.name) if key == "name" else (lambda r: " ".join(
            f"{el}{100 * v:.2f}" for el, v in r.result.composition.items()))
        expect = [r.name for r in sorted(rows, key=val, reverse=desc)]
    else:
        expect = [r.name for r in python_sort(rows, key, desc)]
    assert got == expect
    if key in ("dhmix", "omega"):
        assert got[-1] == "R007", "the N/A (error) row is last in both directions"
    # a page in the middle of the sorted dataset, not a sorted page
    assert [r.name for r in view.page(20, 10)] == expect[20:30]


def test_sort_is_over_the_dataset_not_the_visible_page(repo, engine):
    view = ResultView(repo)
    fill(repo, 3_000, engine)
    view.refresh()
    natural_first_page = {r.id for r in view.page(0, 100)}
    view.request(sort=("dhmix", True))
    top = view.page(0, 100)
    best = repo.con.execute("SELECT max(dhmix) FROM results WHERE err = 0").fetchone()[0]
    assert top[0].raw("dhmix") == best
    assert {r.id for r in top} != natural_first_page
    values = [r.raw("dhmix") for r in view.iterate() if r.raw("dhmix") is not None]
    assert values == sorted(values, reverse=True) and len(values) == 3_000 - 60     # 1 error row per 50


def test_filters_apply_to_the_complete_dataset_with_explicit_na_rules(repo, engine):
    rows = mixed_rows(engine)
    view = ResultView(repo)
    store_rows(repo, rows)
    view.refresh()
    flt = [RangeFilter("delta", None, 6.6), RangeFilter("omega", 1.1, None), ModelFilter("M1", {"SS", "N/A"})]
    view.request(filters=flt)
    expect = [r.name for r in rows if all(f.accepts(r) for f in flt)]
    assert [r.name for r in view.iterate()] == expect and view.count() == len(expect)
    # N/A never passes a range; a missing label counts as "N/A" for a model filter
    view.request(filters=[RangeFilter("dhmix", -1e9, 1e9)])
    assert "R007" not in [r.name for r in view.iterate()]
    view.request(filters=[ModelFilter("M1", {NA})])
    assert "R007" in [r.name for r in view.iterate()]
    sql, params = filter_sql([RangeFilter("dhmix", 1, 2)])
    assert "err = 0" in sql and "IS NOT NULL" in sql and params == (1.0, 2.0)
    assert order_sql("dhmix", True).startswith("(err != 0 OR")


def test_sqlite_stores_nan_as_null_and_keeps_infinities(repo, engine):
    ch = pack_results([{"Fe": 1, "Ni": 1}], [engine.calculate({"Fe": 1, "Ni": 1})])
    ch["num"]["omega"][0] = float("inf")
    ch["num"]["gamma"][0] = float("nan")
    with_metadata(ch, ["X"], "single")
    run = repo.create_run("single")
    rng = repo.append_chunk(run, ch)
    rec = repo.get_result(rng.start)
    assert rec.raw("omega") == float("inf") and rec.raw("gamma") is None


def test_large_build_runs_in_the_background_and_keeps_the_old_page(repo, engine, monkeypatch):
    view = ResultView(repo)
    fill(repo, 2_000, engine)
    view.refresh()
    before = [r.id for r in view.page(0, 10)]
    applied = view.request(sort=("omega", False), wait=0)
    assert not applied and view.pending is not None
    assert [r.id for r in view.page(0, 10)] == before, "the previous view stays until the new one is ready"
    view.wait()
    assert view.pending is None and view.mat is not None
    vals = [r.raw("omega") for r in view.page(0, 50)]
    assert vals == sorted(vals)


def test_superseded_build_is_discarded(repo, engine):
    view = ResultView(repo)
    fill(repo, 5_000, engine)
    view.refresh()
    view.request(sort=("omega", False), wait=0)
    view.request(sort=("dhmix", True), wait=0)
    view.wait()
    assert view.sort_key == "dhmix"
    vals = [r.raw("dhmix") for r in view.page(0, 50)]
    assert vals == sorted(vals, reverse=True)


# =========================================================================== deletion
def test_removing_a_result_deletes_it_from_the_dataset_and_every_view(repo, engine):
    view = ResultView(repo)
    fill(repo, 300, engine)
    view.refresh()
    view.request(sort=("dhmix", False))
    victim = view.page(150, 1)[0]
    neighbours = [r.id for r in view.page(149, 3)]
    assert view.remove([victim]) == 1
    assert repo.get_result(victim.id) is None
    assert view.count() == view.total() == 299
    assert [r.id for r in view.page(149, 2)] == [neighbours[0], neighbours[2]]
    view.request(sort=(None, False))
    assert victim.id not in {r.id for r in view.iterate()} and view.count() == 299
    assert repo.count() == 299


def test_trash_clears_every_result_of_the_table_but_not_new_ones(repo, engine):
    view = ResultView(repo)
    run = fill(repo, 500, engine)
    view.refresh()
    view.request(filters=[RangeFilter("dhmix", None, 0)])
    view.clear()
    assert view.total() == view.count() == 0
    repo.writer.flush()
    assert repo.count() == 0 and repo.get_run(run).cleared
    fill(repo, 20, engine)
    view.refresh()
    view.wait()
    assert view.total() == 20 and view.count() == repo.count([RangeFilter("dhmix", None, 0)])


# =========================================================================== single writer
def test_writer_batches_commits_and_keeps_ids_contiguous(repo, engine):
    run = repo.create_run("sweep")
    for i in range(40):
        assert repo.submit_chunk(run, synthetic_chunk(engine, 250, i * 250))
    repo.writer.flush()
    repo.finalize_run(run, "complete")
    w = repo.writer
    assert w.rows_written == 10_000
    assert w.commits < 40, f"{w.commits} commits for 40 chunks: rows are batched"
    ids = [r[0] for r in repo.con.execute("SELECT result_id FROM results ORDER BY result_id")]
    assert ids == list(range(ids[0], ids[0] + 10_000))
    seqs = [r[0] for r in repo.con.execute("SELECT seq FROM results ORDER BY result_id")]
    assert seqs == list(range(10_000))


def test_wal_readers_never_see_or_wait_for_an_open_write(repo, engine):
    view = ResultView(repo)
    fill(repo, 100, engine)
    view.refresh()
    assert repo.con.execute("PRAGMA journal_mode").fetchone()[0] == "wal"
    other = sqlite3.connect(repo.db_path, isolation_level=None, timeout=1)
    other.execute("BEGIN IMMEDIATE")
    other.execute("DELETE FROM results")
    t = time.perf_counter()
    assert len(view.page(0, 100)) == 100, "uncommitted deletes are invisible"
    assert time.perf_counter() - t < 0.5, "a reader is not blocked by the writer"
    other.execute("ROLLBACK")
    other.close()


def test_cancel_with_a_full_writer_queue_is_fast_and_keeps_a_contiguous_prefix(engine, tmp_path):
    repo = ResultRepository(engine, root=tmp_path / "s")
    svc = CalculationService(workers=2, inprocess_max=10, channel=repo.writer.inbox)
    repo.inject(delay=0.4)                          # a disk far slower than the calculation
    try:
        recipe = SweepRecipe(tuple(SweepSpec(e, 2, 40, 2) for e in ("Fe", "Co", "Ni", "Cr", "Mn")), "at")
        run = repo.create_run("sweep", **{k: v for k, v in recipe.run_meta().items() if k in ("elements", "axes")})
        job = CalculationJob(JobInput("sweep", (), recipe, (), {}), engine, delivery(repo, run, svc))
        svc.start(job)
        t0 = time.perf_counter()
        while job.chunks_emitted < 8 and time.perf_counter() - t0 < 60:
            time.sleep(0.02)
        assert job.chunks_emitted - repo.writer.stored_chunks(run) > 3, "chunks are waiting for the writer"
        t = time.perf_counter()
        job.cancel()
        job.wait(10)
        assert job.state == "cancelled" and time.perf_counter() - t < 1.5
        repo.inject(delay=0)
        res = repo.finalize_run(run, "cancelled", upto=job.chunks_emitted)
        assert res["calculated"] == job.done_count, "exactly the chunks the job handed on are kept"
        info = repo.get_run(run)
        seqs = [r[0] for r in repo.con.execute("SELECT seq FROM results WHERE run_id=? ORDER BY result_id", (run,))]
        assert seqs == list(range(len(seqs))) and 0 < len(seqs) == info.calculated < 94_381
        assert info.status == "cancelled"
    finally:
        svc.shutdown()
        repo.close()


def test_writer_failure_stops_the_run_and_reports_plainly(repo, engine):
    view = ResultView(repo)
    run = repo.create_run("sweep")
    assert repo.submit_chunk(run, synthetic_chunk(engine, 100))
    repo.writer.flush()
    repo.inject(fail="database or disk is full")
    repo.submit_chunk(run, synthetic_chunk(engine, 100, 100))
    repo.writer.flush()
    assert repo.writer.error == "The disk holding the temporary results database is full."
    assert not repo.submit_chunk(run, synthetic_chunk(engine, 10)), "a failed writer takes no more chunks"
    repo.finalize_run(run, "complete")
    info = repo.get_run(run)
    assert info.status == "failed" and info.calculated == 100
    view.refresh()
    assert view.total() == 100, "rows committed before the failure stay valid"


def test_repository_creation_failure_is_reported_without_sqlite_text(tmp_path):
    blocker = tmp_path / "file"
    blocker.write_text("x")
    with pytest.raises(R.RepositoryError) as exc:
        ResultRepository(root=blocker)
    assert "temporary results database" in str(exc.value)


# =========================================================================== lifecycle
def test_close_deletes_the_session_and_stops_its_threads(engine, tmp_path):
    repo = ResultRepository(engine, root=tmp_path / "s")
    fill(repo, 100, engine)
    d = repo.session_dir
    repo.close()
    assert not d.exists()
    names = {t.name for t in threading.enumerate()}
    assert not repo.writer.alive, "the writer process has exited"
    assert not any(t.name in ("hea-view-builder", "hea-db-writer-replies") and t.is_alive()
                   for t in threading.enumerate()), names


def test_stale_sessions_are_removed_and_a_corrupt_one_does_not_matter(engine, tmp_path):
    root = tmp_path / "s"
    stale = root / "session-dead"
    stale.mkdir(parents=True)
    (stale / "results.sqlite").write_bytes(b"not a database at all")
    (stale / "owner.pid").write_text("999999 0")
    repo = ResultRepository(engine, root=root)
    try:
        assert not stale.exists(), "a crashed session's folder is removed at the next start"
        assert repo.count() == 0
        if os.name == "nt":       # a live session's open database cannot be deleted under it
            assert cleanup_stale_sessions(root) == 0 and repo.session_dir.exists()
    finally:
        repo.close()


# =========================================================================== provenance / plots
def run_job(svc, repo, engine, recipe, options=None):
    meta = recipe.run_meta()
    run = repo.create_run(meta["mode"], source=meta["source"], axes=meta["axes"], elements=meta["elements"])
    job = CalculationJob(JobInput(meta["mode"], (), recipe, (), options or {}), engine, delivery(repo, run, svc))
    svc.start(job)
    job.wait(120)
    assert job.state == "done", job.error
    repo.finalize_run(run, "complete", upto=job.chunks_emitted)
    return run, job


@pytest.fixture(scope="module")
def pool():
    svc = CalculationService(workers=2, inprocess_max=20)
    yield svc
    svc.shutdown()


@pytest.mark.parametrize("inprocess", [True, False])
def test_expression_variables_and_blocks_survive_calculation(engine, tmp_path, pool, inprocess):
    repo = ResultRepository(engine, root=tmp_path / "s")
    svc = CalculationService(workers=1) if inprocess else pool
    try:
        text = "(FeCoNi)1-xAlx\n\nCr{y}Fe{z}Ni"
        exprs = parse_blocks(text)
        ranges = {"x": VariableRange("x", 0, 0.5, 0.05), "y": VariableRange("y", 0.1, 0.5, 0.1),
                  "z": VariableRange("z", 0.1, 0.3, 0.1)}
        recipe = ExpressionRecipe(tuple(exprs), tuple(ranges.items()), text)
        view = ResultView(repo)
        run, job = run_job(svc, repo, engine, recipe)
        view.refresh()
        ref = generate_blocks(exprs, ranges)
        recs = list(view.iterate())
        assert [r.name for r in recs] == ref.names, "same order and duplicate removal as the in-memory generator"
        assert [r.block for r in recs] == ref.blocks
        assert [r.variables for r in recs] == ref.variables
        info = repo.get_run(run)
        assert [b["expression"] for b in info.source["blocks"]] == ["(FeCoNi)1-xAlx", "Cr{y}Fe{z}Ni"]
        assert info.source["ranges"]["y"] == [0.1, 0.5, 0.1]
        # future plots: x vs δ (block 0) and a y/z map coloured by Φ and by the M1 label
        series = repo.get_series(run, "x", "delta")
        assert [x for x, _ in series] == sorted(v["x"] for v in ref.variables if "x" in v)
        surf = repo.get_surface(run, "y", "z", "M1", fixed=None)
        assert len(surf) == sum(1 for b in ref.blocks if b == 1)
        fixed = repo.get_series(run, "y", "phi_king", fixed={"z": 0.2})
        assert len(fixed) == len({round(v["y"], 6) for v in ref.variables if "z" in v and abs(v["z"] - 0.2) < 1e-9})
        al = repo.get_series(run, "x", "Al at.%")
        assert all(abs(a - 100 * x) < 1e-9 for x, a in al)
    finally:
        repo.close()


@pytest.mark.parametrize("unit", ["at", "ratio"])
def test_composition_space_coordinates_survive_calculation(engine, tmp_path, pool, unit):
    repo = ResultRepository(engine, root=tmp_path / "s")
    try:
        specs = (SweepSpec("Al", 0, 30, 10, False), SweepSpec("Ni", 20, 60, 10), SweepSpec("Co", 10, 50, 10)) \
            if unit == "at" else (SweepSpec("Al", 0.5, 1.5, 0.5), SweepSpec("Ni", 1, 2, 0.5), SweepSpec("Co", 1, 1, 0))
        recipe = SweepRecipe(specs, unit)
        view = ResultView(repo)
        run, job = run_job(pool, repo, engine, recipe)
        view.refresh()
        _src, names, comps, _env = recipe()
        recs = list(view.iterate())
        assert [r.name for r in recs] == names and job.total == len(comps)
        for rec, comp in zip(recs, comps):
            assert rec.design == {el: float(comp.get(el, 0.0)) for el in ("Al", "Ni", "Co")}
            assert rec.input_amounts == comp
        axes = {a["name"] for a in repo.available_axes(run)}
        assert {"Al", "Ni", "Co", "Al at.%", "Ni at.%", "dhmix", "M1"} <= axes
        surf = repo.get_surface(run, "Al at.%", "Ni at.%", "omega")
        assert len(surf) == len(comps)
        for x, y, _v in surf:
            assert 0 <= x <= 100 and 0 < y <= 100
        grid = repo.aggregate_plot_data(run, "Al at.%", "Ni at.%", "M1", bins=(4, 4), agg="mode")
        assert sum(c[3] for c in grid["cells"]) <= len(comps) and grid["cells"]
        chunks = list(repo.stream_plot_data(run, ["Ni at.%", "dhmix", "M1"], chunk=7))
        assert sum(len(c[0]) for c in chunks) == len(comps)
    finally:
        repo.close()


def test_recalculation_reads_inputs_from_the_database_and_keeps_provenance(engine, tmp_path, pool):
    repo = ResultRepository(engine, root=tmp_path / "s")
    try:
        exprs = parse_blocks("(FeCoNi)1-xAlx")
        recipe = ExpressionRecipe(tuple(exprs), (("x", VariableRange("x", 0, 1, 0.02)),), "(FeCoNi)1-xAlx")
        run, _ = run_job(pool, repo, engine, recipe)
        old = repo.get_run(run)
        view = ResultView(repo)
        view.lo = old.first_id
        view.refresh()
        before = [(r.name, r.variables, r.block) for r in view.iterate()]
        lo, hi = view.hide_all()
        rr = RecalcRecipe(str(repo.db_path), run, lo, hi, tuple(a["col"] for a in old.axes if a.get("col")))
        meta = {"run_meta": None}
        new = repo.create_run(old.mode, axes=old.axes, elements=old.elements, source={"recalculated_from": run})
        job = CalculationJob(JobInput("recalc", (), rr, (), {"annealing_temperature": 900.0}), engine,
                             delivery(repo, new, pool))
        pool.start(job)
        job.wait(120)
        assert job.state == "done", job.error
        repo.finalize_run(new, "complete", upto=job.chunks_emitted)
        view.refresh()
        after = list(view.iterate())
        assert [(r.name, r.variables, r.block) for r in after] == before
        assert all(r.temperature == 900.0 for r in after if r.ok)
        _ = meta
    finally:
        repo.close()


# =========================================================================== export
def test_xlsx_export_splits_worksheets_instead_of_truncating(repo, engine, tmp_path):
    import openpyxl
    from hea_calc.core.export import write_records
    view = ResultView(repo)
    fill(repo, 25, engine)
    view.refresh()
    out = tmp_path / "x.xlsx"
    n = write_records(view.iterate(), out, ["Fe", "Co", "Ni", "Cr", "Mn"], max_sheet_rows=10)
    wb = openpyxl.load_workbook(out, read_only=True)
    assert n == 25
    assert wb.sheetnames[:3] == ["Results", "Results (2)", "Results (3)"]
    rows = sum(ws.max_row - 1 for ws in wb.worksheets[:3])
    assert rows == 25
    wb.close()


def test_export_process_streams_a_consistent_snapshot(repo, engine, tmp_path):
    import multiprocessing as mp
    from hea_calc.core.export import export_process_main
    view = ResultView(repo)
    fill(repo, 1_234, engine)
    view.refresh()
    view.request(filters=[RangeFilter("dhmix", None, -2.0)])
    ctx = mp.get_context("spawn")
    out = tmp_path / "f.csv"
    q, val, ev = ctx.Queue(), ctx.Value("q", 0), ctx.Event()
    p = ctx.Process(target=export_process_main, args=(view.export_spec(), str(out), "file",
                                                     {"elements": ["Fe", "Co"], "notes": {}}, val, ev, q))
    p.start()
    status, n = q.get(timeout=60)
    p.join(10)
    assert status == "ok" and n == view.count() and val.value == n
    with out.open(encoding="utf-8-sig") as fh:
        body = list(csv.reader(fh))[1:]
    assert [r[0] for r in body] == [r.name for r in view.iterate()]


# =========================================================================== UI on the repository
@pytest.fixture(scope="module")
def app():
    a = QApplication.instance() or QApplication([])
    from hea_calc.ui.theme import load_fonts
    load_fonts()
    return a


@pytest.fixture()
def win(app):
    from hea_calc.ui.main_window import MainWindow
    from hea_calc.ui.theme import THEME
    THEME.set_theme("dark")
    w = MainWindow(settings=QSettings(str(Path(tempfile.mkdtemp()) / "s.ini"), QSettings.Format.IniFormat))
    w.resize(1920, 1000)
    w.show()
    app.processEvents()
    yield w
    w.close()
    w.deleteLater()
    app.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    app.processEvents()


def dbl(view, r, c=2):
    vp = view.viewport()
    pos = view.visualRect(view.model().index(r, c)).center()
    QApplication.sendEvent(vp, QMouseEvent(QEvent.Type.MouseButtonDblClick, QPointF(pos), QPointF(vp.mapToGlobal(pos)),
                                           Qt.MouseButton.LeftButton, Qt.MouseButton.LeftButton,
                                           Qt.KeyboardModifier.NoModifier))


def loaded(win, engine, n=600):
    fill(win.repo, n, engine)
    win.table.model.refresh_committed()
    QApplication.processEvents()
    return win.table.model


def test_double_click_after_numeric_sort_opens_the_right_alloy(win, app, engine):
    m = loaded(win, engine)
    m.sort_by("omega")
    m.sort_by("omega")                                              # descending, whole dataset
    m.wait()
    m.set_page(3)
    app.processEvents()
    rec = m.row_at(4)
    dbl(win.table.view, 4)
    app.processEvents()
    assert win.details.row.row_id == rec.id and win.details.row.name == rec.name
    best = win.repo.con.execute(f"SELECT name FROM results WHERE err=0 AND result_id >= {m.view.lo} "
                                "ORDER BY omega DESC, result_id LIMIT 1 OFFSET 304").fetchone()[0]
    assert rec.name == best


def test_double_click_after_filtering_and_the_details_window(win, app, engine, monkeypatch):
    m = loaded(win, engine)
    m.set_filters([RangeFilter("dhmix", None, -3.0)])
    m.wait()
    rec = m.row_at(7)
    assert rec.raw("dhmix") <= -3.0
    dbl(win.table.view, 7)
    app.processEvents()
    assert win.details.row.row_id == rec.id
    from hea_calc.ui.widgets import dialogs
    shown = []
    monkeypatch.setattr(dialogs.AlloyDetailsDialog, "exec", lambda self: shown.append(self.panel.row) or 0)
    menu = win.table.view.row_menu(m.row_at(2), "name")
    next(a for a in menu.actions() if a.text() == "Open details in new window").trigger()
    assert shown[0].row_id == m.row_at(2).id and shown[0].name == m.row_at(2).name


def test_remove_result_on_the_last_page_moves_to_a_valid_page(win, app, engine):
    m = loaded(win, engine, 201)
    win.table.pagination.size_box.setCurrentIndex(3)                 # 100 per page
    m.set_page(2)
    assert m.rowCount() == 1
    victim = m.row_at(0)
    dbl(win.table.view, 0)
    app.processEvents()
    win.table.view._remove(victim)
    app.processEvents()
    assert m.page == 1 and m.rowCount() == 100 and m.total_count() == 200
    assert win.repo.get_result(victim.id) is None
    assert win.details.row is None, "Details of a removed result are cleared"


def test_removing_the_last_filtered_result_shows_the_no_matches_state(win, app, engine):
    m = loaded(win, engine, 100)
    target = m.row_at(17)
    m.set_filters([RangeFilter("dhmix", target.raw("dhmix"), target.raw("dhmix"))])
    m.wait()
    while m.rowCount():
        win.table.view._remove(m.row_at(0))
    app.processEvents()
    assert m.view_count() == 0 and m.total_count() > 0
    assert win.table.empty.isVisible() and win.table.empty.title == "No matching results"


def test_export_all_and_filtered_cover_the_dataset_not_the_page(win, app, engine, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    m = loaded(win, engine, 1_500)
    win.table.pagination.size_box.setCurrentIndex(0)                 # 10 rows visible
    m.set_filters([RangeFilter("dhmix", None, -2.5)])
    m.wait()
    expect_filtered = m.view_count()
    for scope, name in (("all", "all.csv"), ("filtered", "flt.csv")):
        out = tmp_path / name
        monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, o=out, **k: (str(o), ""))
        win.export("csv", scope)
        t0 = time.perf_counter()
        while win.exporter.running and time.perf_counter() - t0 < 60:
            app.processEvents()
            time.sleep(0.01)
        app.processEvents()
        with out.open(encoding="utf-8-sig") as fh:
            n = sum(1 for _ in fh) - 1
        assert n == (1_500 if scope == "all" else expect_filtered), scope
    assert 10 < expect_filtered < 1_500


def test_copy_visible_rows_is_page_scoped_and_copy_table_is_view_scoped(win, app, engine):
    m = loaded(win, engine, 3_000)
    win.table.pagination.size_box.setCurrentIndex(1)                 # 20
    win.copy_visible_rows()
    assert len(QApplication.clipboard().text().strip("\r\n").split("\r\n")) == 21
    QApplication.clipboard().setText("")
    win.copy_table()                                                  # 3 000 rows: background process
    t0 = time.perf_counter()
    while win.exporter.running and time.perf_counter() - t0 < 60:
        app.processEvents()
        time.sleep(0.01)
    app.processEvents()
    assert len(QApplication.clipboard().text().strip("\r\n").split("\r\n")) == 3_001


def test_storage_failure_during_a_calculation_never_leaves_the_ui_calculating(win, app):
    win.repo.inject(fail="database or disk is full")
    for el in ("Fe", "Co", "Ni", "Cr"):
        win.elements.table.tiles[el].click()
    win.composition.set_mode(1)
    win.composition.unit_box.setCurrentIndex(1)
    for r in win.composition.sweep.rows.values():
        r.min.setText("5"), r.max.setText("50"), r.step.setText("5")
    win.composition.calc_btn.click()
    t0 = time.perf_counter()
    while win.worker is not None and time.perf_counter() - t0 < 60:
        app.processEvents()
        time.sleep(0.01)
    app.processEvents()
    assert win.worker is None and win.composition.calc_btn.isEnabled()
    assert "disk" in win.notification.body.text() and "database or disk is full" not in win.notification.body.text()


def test_main_process_memory_does_not_grow_with_the_dataset(win, app, engine):
    """250 000 more rows through the GUI's repository: private bytes stay nearly flat (the former
    columnar store grew 0.39 KiB/row = 95 MiB here; ResultRow objects 1.4 GiB)."""
    if os.name != "nt":
        pytest.skip("private bytes via the Windows API")
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
    from perf_harness import memory_mib
    m = loaded(win, engine, 50_000)
    m.set_page(m.page_count() - 1)
    gc.collect()
    before = memory_mib()["privateMiB"]
    fill(win.repo, 250_000, engine)
    m.refresh_committed()
    m.set_page(m.page_count() - 1)
    app.processEvents()
    gc.collect()
    after = memory_mib()["privateMiB"]
    assert m.total_count() == 300_000 and m.rowCount() <= 100
    assert after - before < 45, (before, after)
    assert len(m.findChildren(QObject)) < 50, "no QObject per result"


# =========================================================================== size tiers (former 200 000 guard)
def test_size_tiers_and_capacity_come_from_the_measured_limits(monkeypatch, tmp_path):
    from hea_calc.core import limits
    assert limits.size_class(249_999) == "normal" and limits.size_class(250_000) == "large"
    assert limits.size_class(1_000_000) == "very_large" and limits.size_class(10_000_001) == "too_large"
    est = limits.estimate(1_000_000, 8)
    assert 40 < est.seconds < 80 and est.bytes == 1_000_000 * limits.BYTES_PER_RESULT
    assert limits.describe(est)[0] == "Very large set" and "stored on disk" in limits.describe(est)[1]
    assert "largest tested" in limits.capacity_problem(10_000_001, None)
    assert limits.capacity_problem(1_000_000, tmp_path) is None
    monkeypatch.setattr(limits.shutil, "disk_usage", lambda _p: type("U", (), {"free": 100 * 2**20})())
    assert "Not enough free disk space" in limits.capacity_problem(1_000_000, tmp_path)


@pytest.mark.parametrize("lo, hi, count, tier", [(1, 30, 263_121, "Large calculation set"), (4, 40, 1_039_776, "Very large set")])
def test_sweep_summary_reports_large_sets_instead_of_refusing_them(win, app, lo, hi, count, tier):
    for el in ("Fe", "Co", "Ni", "Cr", "Mn"):
        win.elements.table.tiles[el].click()
    win.composition.set_mode(1)
    win.composition.unit_box.setCurrentIndex(1)
    sw = win.composition.sweep
    for r in sw.rows.values():
        r.min.setText(str(lo)), r.max.setText(str(hi)), r.step.setText("1")
    sw._update_summary()
    t0 = time.perf_counter()
    while sw._counter.busy and time.perf_counter() - t0 < 30:
        app.processEvents()
    app.processEvents()
    assert sw.summary_data.compositions == count and not sw.summary_data.truncated
    text, tip = sw.summary.text(), sw.summary.toolTip()
    assert f"{count:,} valid compositions" in text and "more than" not in text
    assert tip.startswith(tier) and "Estimated on this computer" in tip
    assert sw.recipe().estimate() == count, "no longer refused at 200 000"


def test_beyond_the_largest_tested_batch_is_refused_plainly(win, app):
    for el in ("Fe", "Co", "Ni", "Cr", "Mn", "Al"):
        win.elements.table.tiles[el].click()
    win.composition.set_mode(1)
    win.composition.unit_box.setCurrentIndex(1)
    sw = win.composition.sweep
    for r in sw.rows.values():
        r.min.setText("1"), r.max.setText("40"), r.step.setText("1")      # 41 659 248 alloys
    sw._update_summary()
    t0 = time.perf_counter()
    while sw._counter.busy and time.perf_counter() - t0 < 30:
        app.processEvents()
    app.processEvents()
    assert sw.summary_data.truncated and "more than 10,000,000" in sw.summary.text()
    win.composition.calc_btn.click()
    app.processEvents()
    assert win.worker is None and "largest tested batch" in win.notification.body.text()


def test_clearing_everything_gives_the_disk_space_back(repo, engine):
    view = ResultView(repo)
    fill(repo, 60_000, engine)
    view.refresh()
    view.request(sort=("dhmix", True), wait=60)
    view.request(sort=("omega", False), wait=60)
    repo.writer.flush()
    size = lambda: sum(repo.size_bytes().values())          # noqa: E731
    big = size()
    view.clear()
    repo.writer.flush(timeout=60)
    done = threading.Event()
    repo.builder.submit(R.ViewTask("drop", table="none", done=done))
    done.wait(30)
    time.sleep(0.2)
    small = size()
    assert big > 20 * 2**20 and small < big / 10, (big, small)
    fill(repo, 10, engine)
    view.refresh()
    assert view.stale and view.wants_resort(running=False), "the kept sort must be rebuilt for new rows"
    view.rebuild(wait=30)
    assert view.total() == view.count() == 10


def test_a_single_calculation_with_a_sort_active_appears_at_once(win, app):
    win.calculate("single", ["A"], [{"Fe": 1, "Ni": 1}], [{}])
    win.calculate("single", ["B"], [{"Co": 1, "Cr": 1}], [{}])
    m = win.table.model
    m.sort_by("name")
    m.sort_by("name")                                              # descending
    win.calculate("single", ["C"], [{"Fe": 1, "Co": 1, "Ni": 1}], [{}])
    app.processEvents()
    assert [r.name for r in m.page_rows()] == ["C", "B", "A"]
    assert win.details.row.name == "C"


def test_closing_during_a_calculation_leaves_no_process_and_no_file(app):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
    from perf_harness import process_alive
    from hea_calc.ui.main_window import MainWindow
    w = MainWindow(settings=QSettings(str(Path(tempfile.mkdtemp()) / "s.ini"), QSettings.Format.IniFormat))
    w.show()
    for el in ("Fe", "Co", "Ni", "Cr", "Mn"):
        w.elements.table.tiles[el].click()
    w.composition.set_mode(1)
    w.composition.unit_box.setCurrentIndex(1)
    for r in w.composition.sweep.rows.values():
        r.min.setText("2"), r.max.setText("40"), r.step.setText("2")
    w.composition.calc_btn.click()
    t0 = time.perf_counter()
    while w.table.model.total_count() < 1000 and time.perf_counter() - t0 < 60:
        app.processEvents()
        time.sleep(0.01)
    pids = list(w.jobs.worker_pids()) + [w.repo.writer.pid]
    session = w.repo.session_dir
    assert w.worker is not None and len(pids) > 1
    w.close()
    deadline = time.perf_counter() + 10
    while any(process_alive(p) for p in pids) and time.perf_counter() < deadline:
        time.sleep(0.05)
    assert not any(process_alive(p) for p in pids), "no orphan worker or writer process"
    assert not session.exists(), "the session database is deleted"
    w.deleteLater()
    app.processEvents()
